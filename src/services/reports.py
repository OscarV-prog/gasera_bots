"""Excel and reporting generation service using openpyxl."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.repositories import get_repository

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=15, bold=True, color="0F172A")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="64748B")
SECTION_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
TOTAL_BORDER = Border(
    top=Side(style="thin", color="0F172A"),
    bottom=Side(style="double", color="0F172A"),
)


def _auto_fit_columns(ws, max_cols: int | None = None) -> None:
    """Adjust column widths automatically based on cell content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        if max_cols and col[0].column > max_cols:
            continue
        for cell in col:
            val = cell.value
            if val is not None:
                lines = str(val).split("\n")
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)


def build_orders_sheet(ws, orders: list[dict[str, Any]]) -> None:
    """Populate a worksheet with order and sales data."""
    ws.title = "Pedidos & Ventas"
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.append(["GAS A TU PUERTA - REPORTE DE PEDIDOS Y VENTAS"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total registros: {len(orders)}"])
    ws.cell(row=2, column=1).font = SUBTITLE_FONT
    ws.append([])

    headers = [
        "Folio",
        "Fecha Registro",
        "Cliente",
        "Teléfono",
        "Dirección de Entrega",
        "Horario Solicitado",
        "Productos",
        "Total (MXN)",
        "Forma de Pago",
        "Chofer Asignado",
        "Unidad",
        "Estado del Pedido",
        "Motivo de Rechazo (si aplica)",
        "Referencias / Notas",
    ]
    ws.append(headers)
    header_row = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    status_labels = {
        "confirmed": "Por Asignar",
        "assigned": "Asignado",
        "in_route": "En Camino",
        "delivered": "Entregado",
        "cancelled": "Cancelado",
        "rejected_by_driver": "Rechazado por Chofer",
        "scheduled": "Programado",
    }

    row_idx = 5
    total_revenue = 0.0
    for o in orders:
        items_str = ", ".join(f"{it.get('quantity', 1)}x {it.get('product_name', '')}" for it in o.get("items", []))
        created_str = o.get("created_at", "")
        if created_str and "T" in created_str:
            try:
                dt = datetime.fromisoformat(created_str.replace("Z", "+00:00"))
                created_str = dt.strftime("%d/%m/%Y %H:%M")
            except Exception:
                pass

        total_amount = float(o.get("total_amount") or 0.0)
        total_revenue += total_amount

        row_data = [
            o.get("id"),
            created_str,
            o.get("customer_name") or "Cliente",
            o.get("customer_phone") or "",
            o.get("delivery_address") or "",
            o.get("delivery_schedule") or "Lo antes posible",
            items_str,
            total_amount,
            o.get("payment_method") or "Efectivo",
            o.get("driver_name") or "Sin Asignar",
            o.get("driver_vehicle") or "",
            status_labels.get(o.get("status"), o.get("status")),
            o.get("rejection_reason") or "",
            o.get("notes") or "",
        ]
        ws.append(row_data)

        # Apply borders and formatting
        for c_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")
            if c_idx == 8:  # Currency column
                c.number_format = '"$"#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in (1, 2, 4, 12):
                c.alignment = Alignment(horizontal="center", vertical="center")

        row_idx += 1

    # Total Row
    total_row = [
        "TOTAL",
        "",
        "",
        "",
        "",
        "",
        f"{len(orders)} pedidos",
        total_revenue,
        "",
        "",
        "",
        "",
        "",
        "",
    ]
    ws.append(total_row)
    for c_idx in range(1, len(total_row) + 1):
        c = ws.cell(row=row_idx, column=c_idx)
        c.font = Font(name="Calibri", size=11, bold=True)
        c.border = TOTAL_BORDER
        if c_idx == 8:
            c.number_format = '"$"#,##0.00'
            c.alignment = Alignment(horizontal="right", vertical="center")

    _auto_fit_columns(ws, len(headers))


def build_rejections_sheet(ws, rejections: list[dict[str, Any]]) -> None:
    """Populate a worksheet with incidents and driver rejection reports."""
    ws.title = "Incidencias & Rechazos"
    ws.views.sheetView[0].showGridLines = True

    ws.append(["BITÁCORA DE INCIDENCIAS Y PEDIDOS RECHAZADOS POR CHOFERES"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total incidencias: {len(rejections)}"])
    ws.cell(row=2, column=1).font = SUBTITLE_FONT
    ws.append([])

    headers = [
        "ID Incidencia",
        "Folio Pedido",
        "Fecha / Hora",
        "Chofer que Rechazó",
        "Teléfono Chofer",
        "Unidad / Placa",
        "Motivo del Rechazo",
        "Cliente",
        "Teléfono Cliente",
        "Dirección de Entrega",
        "Total Pedido (MXN)",
        "Estado de Incidencia",
    ]
    ws.append(headers)
    header_row = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    row_idx = 5
    for r in rejections:
        created_str = r.get("created_at", "")
        if created_str and "T" in created_str:
            try:
                dt = datetime.fromisoformat(created_str.replace("Z", "+00:00"))
                created_str = dt.strftime("%d/%m/%Y %H:%M")
            except Exception:
                pass

        resolved_str = "Resuelto / Reasignado" if r.get("is_resolved") else "Pendiente de Reasignar"
        amount = float(r.get("total_amount") or 0.0)

        row_data = [
            r.get("id"),
            r.get("order_id"),
            created_str,
            r.get("driver_name") or "Chofer",
            r.get("driver_phone") or "",
            r.get("driver_vehicle_plate") or "",
            r.get("reason") or "Sin motivo",
            r.get("customer_name") or "",
            r.get("customer_phone") or "",
            r.get("delivery_address") or "",
            amount,
            resolved_str,
        ]
        ws.append(row_data)

        for c_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")
            if c_idx == 11:
                c.number_format = '"$"#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in (1, 2, 3, 5, 12):
                c.alignment = Alignment(horizontal="center", vertical="center")

        row_idx += 1

    _auto_fit_columns(ws, len(headers))


def build_drivers_sheet(ws, drivers: list[dict[str, Any]], repo) -> None:
    """Populate a worksheet with driver fleet details and metrics."""
    ws.title = "Flota & Choferes"
    ws.views.sheetView[0].showGridLines = True

    ws.append(["DIRECTORIO Y RENDIMIENTO DE FLOTA DE REPARTO"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Choferes activos: {len(drivers)}"])
    ws.cell(row=2, column=1).font = SUBTITLE_FONT
    ws.append([])

    headers = [
        "ID Chofer",
        "Nombre",
        "Teléfono",
        "Telegram ID",
        "Tipo de Vehículo",
        "Placa / Número de Unidad",
        "Zona Asignada",
        "Estado Operativo",
        "Disponibilidad",
        "Pedido Activo en Curso",
        "Pedidos Entregados Históricos",
        "Total Facturado Entregado (MXN)",
    ]
    ws.append(headers)
    header_row = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    row_idx = 5
    for d in drivers:
        d_id = d.get("id")
        # Fetch stats for this driver
        driver_orders = repo.get_orders_by_driver("petroil", d_id, active_only=False) if repo else []
        delivered_orders = [o for o in driver_orders if o.status == "delivered"]
        delivered_count = len(delivered_orders)
        delivered_revenue = sum(o.total_amount for o in delivered_orders)

        op_map = {
            "disponible": "🟢 Disponible (En Turno)",
            "en_entrega": "🟡 En Entrega",
            "fuera_servicio": "🔴 Fuera de Turno",
        }
        op_label = op_map.get(d.get("operational_status"), d.get("operational_status"))
        active_order_str = f"#{d.get('active_order_id')}" if d.get("active_order_id") else "Ninguno"

        row_data = [
            d_id,
            d.get("name"),
            d.get("phone") or "",
            d.get("telegram_user_id") or "Sin vincular",
            (d.get("vehicle_type") or "cilindros").capitalize(),
            d.get("vehicle_plate") or "",
            d.get("zone") or "General",
            op_label,
            "Activo" if d.get("is_available") else "Inactivo",
            active_order_str,
            delivered_count,
            delivered_revenue,
        ]
        ws.append(row_data)

        for c_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")
            if c_idx == 12:
                c.number_format = '"$"#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in (1, 3, 4, 9, 10, 11):
                c.alignment = Alignment(horizontal="center", vertical="center")

        row_idx += 1

    _auto_fit_columns(ws, len(headers))


def build_customers_sheet(ws, customers: list[dict[str, Any]]) -> None:
    """Populate a worksheet with customer directory data."""
    ws.title = "Directorio de Clientes"
    ws.views.sheetView[0].showGridLines = True

    ws.append(["DIRECTORIO DE CLIENTES Y FRECUENCIA DE COMPRA"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total clientes: {len(customers)}"])
    ws.cell(row=2, column=1).font = SUBTITLE_FONT
    ws.append([])

    headers = [
        "ID Cliente",
        "Nombre Completo",
        "Teléfono",
        "Tipo de Cliente",
        "Dirección Principal",
        "Notas de Entrega",
        "Pedidos Realizados",
        "Total Gastado (MXN)",
        "Última Compra",
    ]
    ws.append(headers)
    header_row = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    row_idx = 5
    for c in customers:
        created_str = c.get("last_order_date") or c.get("created_at", "")
        if created_str and "T" in created_str:
            try:
                dt = datetime.fromisoformat(created_str.replace("Z", "+00:00"))
                created_str = dt.strftime("%d/%m/%Y %H:%M")
            except Exception:
                pass

        spent = float(c.get("total_spent") or 0.0)
        frecuente = "Frecuente ⭐" if c.get("is_frequent") else "Regular"

        row_data = [
            c.get("id"),
            c.get("name"),
            c.get("phone") or "",
            frecuente,
            c.get("default_address") or "",
            c.get("address_notes") or "",
            c.get("orders_count") or 0,
            spent,
            created_str,
        ]
        ws.append(row_data)

        for col_i in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_i)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if col_i == 8:
                cell.number_format = '"$"#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_i in (1, 3, 4, 7, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx += 1

    _auto_fit_columns(ws, len(headers))


# -----------------------------------------------------------------------------
# Main Exporter Functions returning BytesIO for FastAPI StreamingResponse
# -----------------------------------------------------------------------------

def generate_master_excel(tenant_id: str = "petroil") -> io.BytesIO:
    """Generate a multi-sheet master workbook containing Orders, Rejections, Drivers and Customers."""
    repo = get_repository()
    orders = repo.get_all_orders_admin(tenant_id, limit=2000)
    rejections = repo.get_order_rejections(tenant_id, unresolved_only=False)
    drivers = repo.get_all_drivers_operational_status(tenant_id)
    customers = repo.get_all_customers_admin(tenant_id)

    wb = Workbook()

    # Sheet 1: Orders
    ws_orders = wb.active
    build_orders_sheet(ws_orders, orders)

    # Sheet 2: Rejections
    ws_rejections = wb.create_sheet()
    build_rejections_sheet(ws_rejections, rejections)

    # Sheet 3: Drivers
    ws_drivers = wb.create_sheet()
    build_drivers_sheet(ws_drivers, drivers, repo)

    # Sheet 4: Customers
    ws_customers = wb.create_sheet()
    build_customers_sheet(ws_customers, customers)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_orders_excel(tenant_id: str = "petroil", status: str | None = None) -> io.BytesIO:
    """Generate an Excel file specifically for orders."""
    repo = get_repository()
    orders = repo.get_all_orders_admin(tenant_id, status=status, limit=3000)
    wb = Workbook()
    ws = wb.active
    build_orders_sheet(ws, orders)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_rejections_excel(tenant_id: str = "petroil") -> io.BytesIO:
    """Generate an Excel file for rejections and incidents."""
    repo = get_repository()
    rejections = repo.get_order_rejections(tenant_id, unresolved_only=False)
    wb = Workbook()
    ws = wb.active
    build_rejections_sheet(ws, rejections)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_drivers_excel(tenant_id: str = "petroil") -> io.BytesIO:
    """Generate an Excel file for fleet and drivers."""
    repo = get_repository()
    drivers = repo.get_all_drivers_operational_status(tenant_id)
    wb = Workbook()
    ws = wb.active
    build_drivers_sheet(ws, drivers, repo)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_customers_excel(tenant_id: str = "petroil") -> io.BytesIO:
    """Generate an Excel file for customers."""
    repo = get_repository()
    customers = repo.get_all_customers_admin(tenant_id)
    wb = Workbook()
    ws = wb.active
    build_customers_sheet(ws, customers)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
